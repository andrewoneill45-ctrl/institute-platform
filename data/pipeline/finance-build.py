#!/usr/bin/env python3
"""Forecast's ledger: every school's finance return, year by year.

Reads the DfE's CFR workbooks (LA-maintained schools) and AAR downloads
(academies) from a folder, harmonises the headline lines across the years
the headers drifted, keeps only full-year (12-month) returns, and writes
apps/platform/public/data/finance.json keyed by URN.

    python3 finance-build.py /path/to/raw apps/platform/public/data/finance.json

Figures are whole pounds in the source (AAR rounded to £1,000 at source);
here they are stored in £ thousands (integers) to keep the file small.
"""
import glob, json, re, sys, os
import openpyxl

RAW, OUT = sys.argv[1], sys.argv[2]

# the lines of the ledger, in the order they are stored per year
FIELDS = ["inc", "exp", "grant", "self", "teach", "supply", "edsup", "admin", "staff", "prem", "energy", "cater", "edsupplies", "prof", "bal", "res", "pupils", "teachers", "fsm", "p16"]
# header matching: (candidates, pick) where pick is 'first' or 'last'
MATCH = {
    "inc":        (["total income"], "last"),
    "exp":        (["total expenditure"], "last"),
    "grant":      (["grant funding", "total grant funding"], "last"),
    "self":       (["self generated funding", "total self generated funding"], "last"),
    "teach":      (["teaching staff"], "first"),
    "supply":     (["supply staff"], "last"),
    "edsup":      (["education support staff"], "first"),
    "admin":      (["administrative and clerical staff"], "first"),
    "staff":      (["staff total", "total staff costs"], "last"),
    "prem":       (["premises:", "premises costs", "premises"], "last"),
    "energy":     (["energy"], "first"),
    "cater":      (["catering exp", "total catering"], "last"),
    "edsupplies": (["educational supplies", "total costs of educational supplies"], "last"),
    "prof":       (["brought in professional", "costs of brought in professional"], "last"),
    "bal":        (["in-year balance", "in year balance"], "last"),
    "res":        (["revenue reserve"], "first"),
    "pupils":     (["no pupils", "number of pupils in academy", "full time equivalent number of pupils in scho", "number of pupils (fte)"], "first"),
    "teachers":   (["no teachers", "fte number of teachers", "number of teachers"], "first"),
    "fsm":        (["eligible for fsm"], "first"),
    "p16":        (["no of pupils in 6th form", "number of pupils in 6th form"], "first"),
    "urn":        (["urn"], "first"),
    "period":     (["period covered", "estimated number of months"], "first"),
    "phase":      (["overall phase"], "first"),
    "name":       (["school name"], "first"),
    "trust":      (["trust or company name", "company name"], "first"),
}
EXCLUDE = {"prem": ["premises staff"], "energy": [], "teach": ["non teaching"], "res": ["share of", "distributed"], "inc": ["balance"], "exp": ["balance"]}

def norm(s):
    return re.sub(r"\s+", " ", str(s or "").replace("\n", " ")).strip().lower()

def find_cols(header):
    h = [norm(x) for x in header]
    cols = {}
    for k, (cands, pick) in MATCH.items():
        hits = []
        for i, x in enumerate(h):
            if k == "urn" and x != "urn":
                continue
            if any(c in x for c in cands) and not any(e in x for e in EXCLUDE.get(k, [])):
                hits.append(i)
        if hits:
            cols[k] = hits[0] if pick == "first" else hits[-1]
    return cols

def year_of(fname):
    m = re.search(r"(20\d\d)-(\d\d)", fname)
    if m:
        return m.group(1) + m.group(2)
    m = re.search(r"SFR(\d+)_(20\d\d)", fname)   # SFR27_2016 is 2014/15; SFR32_2017 is 2015/16
    if m:
        y = int(m.group(2)) - 2
        return f"{y}{str(y + 1)[2:]}"
    return None

def num(v):
    if v is None or v == "":
        return None
    try:
        f = float(str(v).replace(",", ""))
    except ValueError:
        return None
    return f

ledger = {}      # urn -> {year: [fields]}
meta = {}        # urn -> {name, phase, trust?}
years = set()
report = []

def read_sheet(sheet, base, y, src, scale):
    """one data sheet into the ledger; returns (rows seen, rows kept, missing fields)"""
    rows = sheet.iter_rows(values_only=True)
    header = None
    for i, r in enumerate(rows):
        if any(norm(c) == "urn" for c in r):
            header = list(r)
            break
        if i > 8:
            break
    if header is None:
        return None
    cols = find_cols(header)
    # the oldest CFR layouts carry no derived totals: build them from the coded lines
    codes = {}
    for i, h in enumerate(header):
        m = re.match(r"\s*([IE]\d\d[a-z]?)\b", str(h or ""))
        if m and m.group(1) not in codes:
            codes[m.group(1)] = i
    SUMS = {"supply": ["E02", "E10", "E26"], "staff": ["E01", "E02", "E03", "E04", "E05", "E06", "E07", "E08", "E09", "E10", "E11", "E26"],
            "prem": ["E04", "E12", "E13", "E14"], "cater": ["E06", "E25"], "edsupplies": ["E19", "E20", "E21"], "prof": ["E27", "E28"],
            "grant": ["I01", "I02", "I03", "I04", "I05", "I06", "I07", "I15", "I16", "I18"], "self": ["I08", "I09", "I10", "I11", "I12", "I13", "I17"]}
    sums = {k: [codes[c] for c in cs if c in codes] for k, cs in SUMS.items() if k not in cols and sum(c in codes for c in cs) >= len(cs) - 1}
    missing = [k for k in FIELDS if k not in cols and k not in sums]
    n, kept = 0, 0
    for r in rows:
        urn = r[cols["urn"]] if "urn" in cols else None
        try:
            urn = int(float(urn))
        except (TypeError, ValueError):
            continue
        n += 1
        if "period" in cols:
            p = num(r[cols["period"]])
            if p is not None and p < 12:
                continue
        rec = []
        for k in FIELDS:
            if k in cols:
                v = num(r[cols[k]])
            elif k in sums:
                parts = [num(r[i]) for i in sums[k]]
                v = sum(x for x in parts if x is not None) if any(x is not None for x in parts) else None
            else:
                v = None
            if v is None:
                rec.append(None)
            elif k in ("pupils", "teachers", "fsm", "p16"):
                rec.append(round(v, 1))
            else:
                rec.append(int(round(v / scale)))
        if rec[0] is None and rec[1] is None:
            continue
        if "bal" not in cols and rec[0] is not None and rec[1] is not None:
            rec[FIELDS.index("bal")] = rec[0] - rec[1]
        rec.append(src)
        ledger.setdefault(str(urn), {})[y] = rec
        m = meta.setdefault(str(urn), {})
        if "name" in cols and r[cols["name"]]:
            m["n"] = str(r[cols["name"]]).strip()
        if "phase" in cols and r[cols["phase"]]:
            m["ph"] = str(r[cols["phase"]]).strip()
        if "trust" in cols and r[cols["trust"]] and src == "A":
            m["t"] = str(r[cols["trust"]]).strip()
        kept += 1
    return n, kept, missing

for f in sorted(glob.glob(os.path.join(RAW, "*.xlsx"))):
    base = os.path.basename(f)
    y = year_of(base)
    src = "A" if base.upper().startswith(("AAR", "SFR")) else "C"
    if not y:
        continue
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    sfr = base.upper().startswith("SFR")            # the 2014/15 and 2015/16 academies tables: already in £000, one raw sheet per trust type
    scale = 1.0 if sfr else 1000.0
    if sfr:
        sheets = [ws for ws in wb.worksheets if ws.title.startswith("Raw Data") and "Totals" not in ws.title and "CS" not in ws.title]
    else:
        sheets = [ws for ws in wb.worksheets if ws.max_row > 1000][:1]
    if not sheets:
        report.append(f"{base}: no data sheet"); continue
    for sheet in sheets:
        got = read_sheet(sheet, base, y, src, scale)
        if got is None:
            report.append(f"{base} [{sheet.title}]: no header"); continue
        n, kept, missing = got
        years.add(y)
        report.append(f"{base} [{sheet.title}]: {y} {src} rows {n} kept {kept} missing {missing}")

for line in report:
    print(line)

out = {
    "vintage": "DfE consistent financial reporting (CFR, LA-maintained) and academies accounting returns (AAR), " + ", ".join(sorted(years)[:1]) + " to " + ", ".join(sorted(years)[-1:]) + "; full-year returns only; £ thousands",
    "fields": FIELDS + ["src"],
    "years": sorted(years),
    "s": {u: {"m": meta.get(u, {}), "y": ys} for u, ys in ledger.items()},
}
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w") as fh:
    json.dump(out, fh, separators=(",", ":"))
print("schools", len(ledger), "years", sorted(years), "->", OUT, round(os.path.getsize(OUT) / 1e6, 1), "MB")
